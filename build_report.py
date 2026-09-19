#!/usr/bin/env python3
"""Генерация отчёта ЛР1 в формате DOCX."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path(__file__).resolve().parent / "OTCHET_LR1.docx"


def set_run_font(run, name="Times New Roman", size=14, bold=False, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_para(doc, text="", *, size=14, bold=False, italic=False, center=False, space_after=6):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.first_line_indent = Cm(1.25) if text and not center and not bold else Cm(0)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def add_heading_custom(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.first_line_indent = Cm(0)
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text.upper())
        set_run_font(run, size=14, bold=True)
    else:
        run = p.add_run(text)
        set_run_font(run, size=14, bold=True)
    return p


def add_formula(doc, text):
    """Формула по центру, без абзацного отступа."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run(text)
    set_run_font(run, name="Cambria Math", size=13, italic=True)
    # fallback font also on eastAsia
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    return p


def add_matrix_table(doc, title, rows):
    """Матрица как таблица Word — без искажений превью Markdown."""
    if title:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run(title)
        set_run_font(run, size=13, italic=True)

    n_cols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=n_cols)
    table.style = "Table Grid"
    table.autofit = True
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(val))
            set_run_font(run, name="Cambria Math", size=12, italic=True)
    # небольшой отступ после таблицы
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        p.clear()
        run = p.add_run(item)
        set_run_font(run, size=14)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        p.clear()
        run = p.add_run(item)
        set_run_font(run, size=14)


def build():
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)

    # --- Титульная часть ---
    add_para(doc, "Лабораторная работа № 1", center=True, bold=True, size=16, space_after=6)
    add_para(
        doc,
        "Реализация алгоритмов построения проекций трёхмерных объектов.",
        center=True,
        bold=True,
        size=14,
        space_after=2,
    )
    add_para(
        doc,
        "Аффинные преобразования в пространстве",
        center=True,
        bold=True,
        size=14,
        space_after=12,
    )
    add_para(doc, "Объект: объёмная каркасная буква «А»", center=True, size=14, space_after=6)
    add_para(
        doc,
        "Язык: Python 3. Матрицы и проекции реализованы вручную "
        "(без OpenGL / NumPy / готовых 3D-библиотек). "
        "Для окна используется стандартный модуль tkinter.",
        center=True,
        size=12,
        space_after=18,
    )

    # --- 1 ---
    add_heading_custom(doc, "1. Цель работы", 1)
    add_para(
        doc,
        "Разработать программу для выполнения аффинных преобразований и построения "
        "проекций трёхмерного каркасного объекта. Интерфейс должен позволять управлять "
        "текущим преобразованием с помощью мыши или клавиатуры. Необходимо реализовать "
        "элементарные преобразования (перенос, поворот, масштабирование) и дополнительное "
        "динамическое преобразование (анимацию).",
    )

    # --- 2 ---
    add_heading_custom(doc, "2. Каркасная модель", 1)
    add_para(
        doc,
        "Объект задаётся множеством вершин — точек (x, y, z) — и рёбер (пар индексов вершин). "
        "На экране рисуются отрезки между спроецированными вершинами.",
    )
    add_para(
        doc,
        "Буква «А» представлена как призма: передняя и задняя грани связаны рёбрами глубины. "
        "В модели есть вырез между ножками, горизонтальная перекладина и отверстие над ней.",
    )
    add_para(
        doc,
        "Каркасная модель достаточна для изучения преобразований и проекций; "
        "удаление невидимых граней и закраска в данную работу не входят.",
    )

    # --- 3 ---
    add_heading_custom(doc, "3. Однородные координаты и матрицы 4×4", 1)
    add_para(
        doc,
        "Точку (x, y, z) записывают в однородных координатах как столбец из четырёх чисел:",
    )
    add_matrix_table(
        doc,
        "p =",
        [["x"], ["y"], ["z"], ["1"]],
    )
    add_formula(doc, "p = (x, y, z, 1)ᵀ")
    add_para(
        doc,
        "Четвёртая координата нужна для того, чтобы перенос, поворот и масштабирование "
        "записывались одной матрицей M размера 4×4 и комбинировались умножением матриц.",
    )
    add_para(doc, "Преобразование точки:")
    add_formula(doc, "p′ = M · p")
    add_para(doc, "В развёрнутом виде:")
    add_formula(doc, "x′ = m₁₁x + m₁₂y + m₁₃z + m₁₄")
    add_formula(doc, "y′ = m₂₁x + m₂₂y + m₂₃z + m₂₄")
    add_formula(doc, "z′ = m₃₁x + m₃₂y + m₃₃z + m₃₄")
    add_formula(doc, "w′ = m₄₁x + m₄₂y + m₄₃z + m₄₄")
    add_para(
        doc,
        "Для обычных аффинных преобразований w′ = 1. Если w′ ≠ 0 (это важно для перспективы), "
        "декартовы координаты восстанавливают делением:",
    )
    add_formula(doc, "X = x′ / w′,    Y = y′ / w′,    Z = z′ / w′")

    # --- 4 ---
    add_heading_custom(doc, "4. Элементарные аффинные преобразования", 1)
    add_para(
        doc,
        "Аффинное преобразование сохраняет прямолинейность отрезков и параллельность прямых. "
        "В трёхмерном пространстве базовыми являются перенос, поворот и масштабирование.",
    )

    add_heading_custom(doc, "4.1. Перенос (translation)", 2)
    add_matrix_table(
        doc,
        "T =",
        [
            ["1", "0", "0", "dx"],
            ["0", "1", "0", "dy"],
            ["0", "0", "1", "dz"],
            ["0", "0", "0", "1"],
        ],
    )
    add_formula(doc, "x′ = x + dx,    y′ = y + dy,    z′ = z + dz,    w′ = 1")
    add_para(doc, "В программе: клавиши WASD / стрелки, Q/E (перенос по глубине).")

    add_heading_custom(doc, "4.2. Масштабирование (scaling)", 2)
    add_matrix_table(
        doc,
        "S =",
        [
            ["sx", "0", "0", "0"],
            ["0", "sy", "0", "0"],
            ["0", "0", "sz", "0"],
            ["0", "0", "0", "1"],
        ],
    )
    add_formula(doc, "x′ = sx · x,    y′ = sy · y,    z′ = sz · z,    w′ = 1")
    add_para(
        doc,
        "При sx = sy = sz масштаб равномерный (пропорции не искажаются). "
        "Масштабирование выполняется относительно начала координат; модель заранее центрирована. "
        "В программе: колесо мыши, клавиши + / −.",
    )

    add_heading_custom(doc, "4.3. Поворот (rotation)", 2)
    add_para(doc, "Угол φ, c = cos φ, s = sin φ.")

    add_para(doc, "Вокруг оси X:", bold=True)
    add_matrix_table(
        doc,
        "Rx =",
        [
            ["1", "0", "0", "0"],
            ["0", "c", "−s", "0"],
            ["0", "s", "c", "0"],
            ["0", "0", "0", "1"],
        ],
    )
    add_formula(doc, "x′ = x,    y′ = c·y − s·z,    z′ = s·y + c·z")

    add_para(doc, "Вокруг оси Y:", bold=True)
    add_matrix_table(
        doc,
        "Ry =",
        [
            ["c", "0", "s", "0"],
            ["0", "1", "0", "0"],
            ["−s", "0", "c", "0"],
            ["0", "0", "0", "1"],
        ],
    )
    add_formula(doc, "x′ = c·x + s·z,    y′ = y,    z′ = −s·x + c·z")

    add_para(doc, "Вокруг оси Z:", bold=True)
    add_matrix_table(
        doc,
        "Rz =",
        [
            ["c", "−s", "0", "0"],
            ["s", "c", "0", "0"],
            ["0", "0", "1", "0"],
            ["0", "0", "0", "1"],
        ],
    )
    add_formula(doc, "x′ = c·x − s·y,    y′ = s·x + c·y,    z′ = z")
    add_para(
        doc,
        "В программе: R/F, T/G, Y/H; мышь (ЛКМ + движение) изменяет углы вокруг X и Y.",
    )

    # --- 5 ---
    add_heading_custom(doc, "5. Композиция преобразований", 1)
    add_para(
        doc,
        "Несколько преобразований объединяются произведением матриц. Порядок важен: "
        "произведение M₂M₁ означает «сначала M₁, затем M₂» (вектор-столбец справа).",
    )
    add_para(doc, "В программе модельная матрица собирается в порядке:")
    add_numbered(
        doc,
        [
            "масштаб S;",
            "повороты Rx, Ry, Rz;",
            "перенос T.",
        ],
    )
    add_formula(doc, "M = T · Rz · Ry · Rx · S")
    add_para(doc, "С учётом анимации (динамического преобразования) A:")
    add_formula(doc, "Mworld = A · M")

    # --- 6 ---
    add_heading_custom(doc, "6. Проекции трёхмерных объектов", 1)
    add_para(
        doc,
        "Экран плоский, поэтому точки пространства необходимо спроецировать на плоскость.",
    )

    add_heading_custom(doc, "6.1. Ортогональная (параллельная) проекция", 2)
    add_para(
        doc,
        "Проецирующие лучи параллельны. Видимый размер не уменьшается с расстоянием. "
        "В программе используется упрощённая аксонометрия: координата глубины y "
        "подмешивается в экранные координаты, чтобы сохранить ощущение объёма.",
    )
    add_formula(doc, "Xэкр = W/2 + x − 0,45·y")
    add_formula(doc, "Yэкр = H/2 − (z + 0,35·y)")
    add_para(
        doc,
        "Знак у Yэкр выбран так потому, что в окне экранная ось Y направлена вниз. "
        "Переключение: клавиша 1.",
    )

    add_heading_custom(doc, "6.2. Перспективная (центральная) проекция", 2)
    add_para(
        doc,
        "Лучи сходятся в центре проекции (камера). Более далёкие объекты выглядят меньше. "
        "Пусть f — параметр «фокуса», глубина d = y + f. Тогда:",
    )
    add_formula(doc, "xp = (x · f) / d,    zp = (z · f) / d")
    add_formula(doc, "Xэкр = W/2 + xp,    Yэкр = H/2 − zp")
    add_para(doc, "Переключение: клавиша 2.")
    add_para(
        doc,
        "Отличие: при ортогональной проекции размеры не зависят от расстояния; "
        "при перспективной — чем дальше объект, тем он меньше на экране.",
    )

    # --- 7 ---
    add_heading_custom(doc, "7. Конвейер отрисовки кадра", 1)
    add_numbered(
        doc,
        [
            "Взять исходные вершины буквы «А».",
            "Умножить на матрицу Mworld = A · M.",
            "Спроецировать каждую вершину в экранные координаты (X, Y).",
            "Нарисовать рёбра отрезками.",
            "Повторять цикл при включённой анимации.",
        ],
    )
    add_para(
        doc,
        "Важно: готовые 3D-библиотеки не используются. Модуль tkinter только отображает "
        "уже вычисленные двумерные точки и отрезки.",
    )

    # --- 8 ---
    add_heading_custom(doc, "8. Динамическое преобразование", 1)
    add_para(
        doc,
        "По заданию реализовано дополнительное динамическое преобразование (анимация). "
        "По клавише «Пробел» включается/выключается:",
    )
    add_bullets(
        doc,
        [
            "непрерывный поворот вокруг оси Z;",
            "небольшая пульсация масштаба: 1 + 0,08·sin(t).",
        ],
    )
    add_para(
        doc,
        "Это те же аффинные преобразования, но их параметры меняются со временем.",
    )

    # --- 9 ---
    add_heading_custom(doc, "9. Система координат сцены", 1)
    add_bullets(
        doc,
        [
            "X — вправо;",
            "Y — глубина (от наблюдателя);",
            "Z — вверх.",
        ],
    )
    add_para(doc, "На экране дополнительно рисуются оси координат для ориентира.")

    # --- 10 ---
    add_heading_custom(doc, "10. Структура программы", 1)
    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    data = [
        ("Файл", "Назначение"),
        ("letter_a.py", "Вершины и рёбра буквы «А», оси координат"),
        ("math3d.py", "Матрицы 4×4, перенос, поворот, масштаб, проекция"),
        ("main.py", "Окно, ввод пользователя, цикл отрисовки, анимация"),
    ]
    for i, (a, b) in enumerate(data):
        for j, val in enumerate((a, b)):
            cell = table.cell(i, j)
            cell.text = ""
            run = cell.paragraphs[0].add_run(val)
            set_run_font(run, size=12, bold=(i == 0))

    add_para(doc, "Запуск:", space_after=4)
    add_para(doc, "/usr/local/bin/python3 main.py", center=True, size=12)
    add_para(doc, "или ./run.sh", center=True, size=12)

    # --- 11 ---
    add_heading_custom(doc, "11. Управление программой", 1)
    ctrl = [
        ("Действие", "Управление"),
        ("Вращение", "ЛКМ + движение; R/F, T/G, Y/H"),
        ("Перенос", "WASD / стрелки; Q/E"),
        ("Масштаб", "Колесо мыши; + / −"),
        ("Анимация", "Пробел"),
        ("Проекция", "1 — ортогональная, 2 — перспективная"),
        ("Сброс", "0"),
        ("Выход", "Esc"),
    ]
    t2 = doc.add_table(rows=len(ctrl), cols=2)
    t2.style = "Table Grid"
    for i, (a, b) in enumerate(ctrl):
        for j, val in enumerate((a, b)):
            cell = t2.cell(i, j)
            cell.text = ""
            run = cell.paragraphs[0].add_run(val)
            set_run_font(run, size=12, bold=(i == 0))

    # --- 12 ---
    add_heading_custom(doc, "12. Вопросы к защите (краткие ответы)", 1)

    qa = [
        (
            "Что такое аффинное преобразование?",
            "Линейное преобразование плюс перенос; сохраняет прямые и параллельность.",
        ),
        (
            "Зачем однородные координаты?",
            "Чтобы перенос записывался матрицей и все преобразования умножались единообразно: p′ = M·p.",
        ),
        (
            "Чем поворот отличается от масштаба?",
            "Поворот меняет ориентацию объекта, масштаб — его размеры вдоль осей.",
        ),
        (
            "Почему важен порядок умножения матриц?",
            "Произведения TR и RT дают разные результаты.",
        ),
        (
            "Что такое проекция?",
            "Отображение точек трёхмерного пространства на плоскость экрана.",
        ),
        (
            "Чем перспектива отличается от ортогональной проекции?",
            "В перспективе далёкие объекты меньше; в ортогональной размеры от расстояния не зависят.",
        ),
        (
            "Как задана буква «А»?",
            "Списком вершин и рёбер каркаса (призма с вырезами).",
        ),
        (
            "Где в коде реализованы матрицы?",
            "В файле math3d.py: translate, scale, rotate_x/y/z, mat_mul, perspective_project.",
        ),
    ]
    for q, a in qa:
        add_para(doc, q, bold=True, space_after=2)
        # без красной строки для ответа
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        run = p.add_run(a)
        set_run_font(run, size=14)

    # --- 13 ---
    add_heading_custom(doc, "13. Заключение", 1)
    add_para(
        doc,
        "В ходе лабораторной работы реализована программа построения проекций "
        "объёмной каркасной буквы «А» и выполнения аффинных преобразований в пространстве. "
        "Реализованы перенос, поворот, масштабирование, перспективная и ортогональная "
        "проекции, а также динамическое преобразование. Математика преобразований "
        "выполнена самостоятельно на матрицах 4×4.",
    )

    doc.save(OUT)
    print(f"Saved: {OUT}")


if __name__ == "__main__":
    build()
